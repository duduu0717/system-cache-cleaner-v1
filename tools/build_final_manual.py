from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "系统缓存清理工具V1.0说明书_排版模板.docx"
OUTPUT = ROOT / "docs" / "系统缓存清理工具V1.0说明书.docx"


TEXTS = [
    "系统缓存清理工具v1.0，是一个适用于windows系统的缓存清理工具，这个软件可以扫描用户的临时文件，缩略图缓存，并且显示文件的数量和占用空间，需要用户确认之后执行删除，该删除不可恢复，需谨慎处理，用于清理空间和释放缓存",
    "开发语言用的是Python，界面使用的是 Tkinter。文件扫描、空间统计和清理的功能主要是通过 Python 标准库来实现的，扫描和清理任务在后台运行，不会造成界面卡顿。软件最后使用 PyInstaller打包成 一个exe文件，无需安装 Python 就能在电脑中运行。",
    "解压之后，在程序文件夹下双击“系统缓存清理工具V1.0.exe” 这个文件就可以运行",
    "点击“开始扫描”按钮之后，软件就会开始扫描系统中的缓存文件。",
    "扫描过程中可以看到当前的进度、已经扫描的文件数量和文件的占用空间。",
    "在扫描结果页面可以选择你需要清理的缓存文件。",
    "点击“清理选中项目”之后，需要再次确认清理的项目，之后执行清理，这一步要谨慎检查之后操作。",
    "软件开始清理选中的缓存文件，并且显示当前处理进度。",
    "清理完成后查看清理之后释放的空间、成功清理文件的数量、失败清理的文件数量和清理时间。",
    "打开清理日志之后，可以查看每个文件的处理结果以及失败原因。",
    "清理操作步骤需要用户的确认，文件删除后无法恢复，请在仔细检查需要清理的项目之后再执行。这个软件只会处理用户的临时文件、Windows的临时文件和缩略图缓存，不会额外清理其他的内容。",
]


def set_text(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(7)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


doc = Document(TEMPLATE)
placeholders = [p for p in doc.paragraphs if p.text.startswith("【待填写：")]
if len(placeholders) != 21:
    raise RuntimeError(f"模板占位段落数量不正确：{len(placeholders)}")

text_indexes = [0, 1, 2, 4, 6, 8, 10, 12, 14, 16, 18]
if len(TEXTS) != len(text_indexes):
    raise RuntimeError("正文数量与模板位置不一致")
for index, text in zip(text_indexes, TEXTS):
    set_text(placeholders[index], text)
for index, paragraph in enumerate(placeholders):
    if index not in text_indexes:
        remove_paragraph(paragraph)

# 最终说明书不包含“软件处理范围”和“退出和卸载”子标题，处理范围已写入注意事项。
for paragraph in list(doc.paragraphs):
    if paragraph.text.strip() in {"1. 软件处理范围", "2. 退出和卸载"}:
        remove_paragraph(paragraph)

doc.core_properties.title = "系统缓存清理工具 V1.0 操作说明"
doc.core_properties.subject = "系统缓存清理工具软件操作说明书"
doc.save(OUTPUT)
print(OUTPUT)
