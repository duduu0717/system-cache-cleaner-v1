from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
IMAGES = ROOT / "docs" / "manual_images"
OUTPUT = ROOT / "docs" / "系统缓存清理工具V1.0说明书_排版模板.docx"


def font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def border(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        borders.append(item)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "777777")
    borders.append(bottom)
    props.append(borders)


def page_number(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            item = OxmlElement("w:fldChar")
            item.set(qn("w:fldCharType"), kind)
        else:
            item = OxmlElement("w:instrText" if text == " PAGE " else "w:t")
            item.text = text
        run._r.append(item)
    font(run, size=9)


def setup(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.65)
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Cm(17))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    border(left); border(right)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    font(left.paragraphs[0].add_run("系统缓存清理工具 V1.0"), size=9)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    page_number(right.paragraphs[0])
    blank = header.paragraphs[0]
    blank._element.getparent().remove(blank._element)


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(26)
    font(p.add_run(text), "黑体", 18, True)


def heading(doc, text, sub=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(text), "黑体", 11 if sub else 13, True)


def placeholder(doc, prompt, lines=3):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run(f"【待填写：{prompt}】"), "宋体", 10.5, False, "777777")
    for _ in range(lines - 1):
        p.add_run("\n")
        font(p.add_run("________________________________________________________________________________"), "宋体", 9, False, "BBBBBB")


def image(doc, filename, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(IMAGES / filename), width=Inches(6.25))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run(caption), size=9)


def new_page(doc):
    doc.add_page_break()


doc = Document()
setup(doc.sections[0])
doc.styles["Normal"].font.name = "宋体"
doc.styles["Normal"]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
doc.styles["Normal"].font.size = Pt(10.5)

# 第 1 页：简介和技术组成
title(doc, "系统缓存清理工具操作说明")
heading(doc, "一、软件简介")
placeholder(doc, "用 120～180 字说明软件用途、适用对象和目前支持的缓存类别。", 5)
heading(doc, "二、技术组成")
placeholder(doc, "用 100～150 字说明 Python、Tkinter、文件扫描、后台线程和 EXE 打包方式。", 5)

# 第 2 页：启动
new_page(doc)
heading(doc, "三、如何启动")
placeholder(doc, "用 80～120 字说明解压、双击 EXE，以及何时需要管理员权限。", 3)
image(doc, "01_package.png", "图 1  软件启动文件位置")
placeholder(doc, "说明红框内文件的作用，以及首次启动可能出现的系统提示。", 2)

# 第 3 至 9 页：常用操作
operations = [
    ("1. 开始扫描", "02_home.png", "图 2  软件首页", "说明首页信息和“开始扫描”按钮，80～120 字。"),
    ("2. 等待扫描完成", "03_scan.png", "图 3  缓存扫描页面", "说明扫描过程、进度显示和等待要求，60～100 字。"),
    ("3. 选择清理项目", "04_results_real.png", "图 4  扫描结果页面", "说明复选框、文件数量、占用空间和清理按钮，100～150 字。"),
    ("4. 确认清理", "05_confirm_real.png", "图 5  清理确认窗口", "说明如何核对项目，点击“是”或“否”分别会发生什么，80～120 字。"),
    ("5. 执行清理", "06_clean.png", "图 6  缓存清理页面", "说明清理进度、文件占用和权限不足时的处理，80～120 字。"),
    ("6. 查看清理结果", "07_complete_verification.png", "图 7  清理完成页面", "说明释放空间、成功数、失败数、耗时和返回首页按钮，100～150 字。"),
    ("7. 查看清理日志", "08_log.png", "图 8  清理日志内容", "说明日志位置、记录内容和对外发送前的检查，80～120 字。"),
]
for name, filename, caption, prompt in operations:
    new_page(doc)
    if name.startswith("1."):
        heading(doc, "四、常用操作")
    heading(doc, name, sub=True)
    placeholder(doc, prompt, 2)
    image(doc, filename, caption)
    placeholder(doc, "结合截图中的红框，写 1～2 句具体操作。", 2)

# 第 10 页：注意事项
new_page(doc)
heading(doc, "五、注意事项")
placeholder(doc, "说明清理前保存工作、文件删除后无法恢复、系统目录权限和文件占用问题，150～220 字。", 6)
heading(doc, "1. 软件处理范围", sub=True)
placeholder(doc, "说明软件不会处理桌面、文档、下载、浏览器密码和回收站等内容，100～150 字。", 4)
heading(doc, "2. 退出和卸载", sub=True)
placeholder(doc, "说明清理过程中不要关闭软件，以及免安装软件的卸载方法，80～120 字。", 3)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
