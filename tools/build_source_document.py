from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "系统缓存清理工具V1.0源代码.docx"
SOURCE_FILES = [
    ROOT / "run_app.py",
    ROOT / "src" / "cache_cleaner" / "__init__.py",
    ROOT / "src" / "cache_cleaner" / "__main__.py",
    ROOT / "src" / "cache_cleaner" / "core.py",
    ROOT / "src" / "cache_cleaner" / "app.py",
]
MIN_LINES_PER_PAGE = 50


def set_font(run, name="Consolas", size=8.0, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_header_border(cell):
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        borders.append(item)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "777777")
    borders.append(bottom)
    properties.append(borders)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, value, end):
        run._r.append(element)
    set_font(run, "宋体", 9)


def configure_page(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.15)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.55)

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Cm(18.5))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    set_header_border(left)
    set_header_border(right)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_font(left.paragraphs[0].add_run("系统缓存清理工具 V1.0"), "宋体", 9)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(right.paragraphs[0])
    blank = header.paragraphs[0]
    blank._element.getparent().remove(blank._element)


def shade_paragraph(paragraph, color="EEF1F4"):
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def add_file_label(doc, relative_path):
    paragraph = doc.add_paragraph(style="SourceFile")
    paragraph.paragraph_format.keep_with_next = True
    shade_paragraph(paragraph)
    set_font(paragraph.add_run(f"文件：{relative_path}"), "宋体", 8.2, True, "333333")


def add_code_line(doc, line_number, text):
    paragraph = doc.add_paragraph(style="SourceCode")
    number = paragraph.add_run(f"{line_number:>4}  ")
    set_font(number, "Consolas", 8.0, False, "777777")
    code = paragraph.add_run(text if text else " ")
    set_font(code, "Consolas", 8.0, False, "111111")


doc = Document()
configure_page(doc.sections[0])

styles = doc.styles
code_style = styles.add_style("SourceCode", WD_STYLE_TYPE.PARAGRAPH)
code_style.paragraph_format.space_before = Pt(0)
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
code_style.paragraph_format.line_spacing = Pt(11.2)
code_style.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
code_style.font.name = "Consolas"
code_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
code_style.font.size = Pt(8.0)

file_style = styles.add_style("SourceFile", WD_STYLE_TYPE.PARAGRAPH)
file_style.paragraph_format.space_before = Pt(2)
file_style.paragraph_format.space_after = Pt(2)
file_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
file_style.paragraph_format.line_spacing = Pt(10)

all_lines = []
for source_path in SOURCE_FILES:
    relative = source_path.relative_to(ROOT).as_posix()
    lines = source_path.read_text(encoding="utf-8").splitlines()
    for index, text in enumerate(lines):
        all_lines.append((relative if index == 0 else None, text))

page_count = max(1, len(all_lines) // MIN_LINES_PER_PAGE)
base_size, extra_pages = divmod(len(all_lines), page_count)
page_sizes = [base_size + (1 if index < extra_pages else 0) for index in range(page_count)]
page_ends = []
running_total = 0
for size in page_sizes[:-1]:
    running_total += size
    page_ends.append(running_total)

for global_index, (file_label, text) in enumerate(all_lines, start=1):
    if global_index - 1 in page_ends:
        doc.add_page_break()
    if file_label:
        add_file_label(doc, file_label)
    add_code_line(doc, global_index, text)

doc.core_properties.title = "系统缓存清理工具 V1.0 源代码"
doc.core_properties.subject = "软件源代码文档"
doc.core_properties.comments = "完整收录软件生产源码，连续行号排版。"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(f"已生成：{OUTPUT}")
print(f"源代码行数：{len(all_lines)}")
print(f"计划页数：{page_count}")
print(f"每页行数：{page_sizes}")
