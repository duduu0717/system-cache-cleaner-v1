from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "系统缓存清理工具V1.0_操作说明书.docx"
BLUE = "1769E0"
LIGHT = "EAF2FC"


def set_font(run, size=10, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color):
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(item)


def set_cell(cell, text, bold=False, color=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    set_font(run, 9, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, title in zip(table.rows[0].cells, headers):
        shade(cell, BLUE)
        set_cell(cell, title, True, "FFFFFF")
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell(cell, str(value))
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    return table


def add_picture(doc, filename, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "screenshots" / filename), width=Inches(6.55))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(caption), 9, True, "475569")


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10)
for name in ("Heading 1", "Heading 2", "Heading 3"):
    styles[name].font.name = "Microsoft YaHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles[name].font.color.rgb = RGBColor.from_string("173B67")

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("系统缓存清理工具 V1.0  |  操作说明书"), 8, False, "64748B")
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("系统缓存清理工具 V1.0 · 软件操作说明书"), 8, False, "64748B")

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("系统缓存清理工具 V1.0"), 26, True, BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("软件操作说明书"), 20, True, "172033")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("面试测试题交付版"), 11, False, "64748B")
doc.add_paragraph("")
add_table(doc, ["文档属性", "内容"], [("文档版本", "V1.0"), ("软件版本", "V1.0.0"), ("适用平台", "Windows 10 / Windows 11"), ("发布日期", "2026-09-08"), ("文档状态", "正式初版")], [1.6, 4.8])
doc.add_page_break()

doc.add_heading("1. 文档概述", 1)
doc.add_paragraph("本说明书用于指导用户安装并使用“系统缓存清理工具 V1.0”。软件通过扫描预设的 Windows 缓存位置，帮助用户查看缓存占用情况，并在用户明确确认后删除选中的缓存文件。")
doc.add_heading("1.1 核心流程", 2)
doc.add_paragraph("启动软件 → 查看磁盘空间 → 扫描缓存 → 选择清理项目 → 确认清理 → 查看清理结果")
doc.add_heading("1.2 功能范围", 2)
add_table(doc, ["项目", "说明", "默认状态"], [("用户临时文件", "当前 Windows 用户及应用产生的临时文件", "选中"), ("Windows 临时文件", "系统临时目录中当前可访问的文件", "选中"), ("缩略图缓存", "资源管理器生成的缩略图数据库", "选中")], [1.5, 3.8, 1.1])
doc.add_paragraph("V1.0 不包含浏览器历史、密码、回收站、下载目录、大文件查找、重复文件查找和定时清理。")

doc.add_heading("2. 运行环境与启动", 1)
add_table(doc, ["项目", "最低要求"], [("操作系统", "Windows 10 或 Windows 11（64 位）"), ("运行方式", "解压 ZIP 后双击 系统缓存清理工具V1.0.exe"), ("网络", "无需联网"), ("权限", "普通权限可扫描用户缓存；系统目录建议管理员权限")], [1.5, 4.9])
doc.add_heading("2.1 启动步骤", 2)
for text in ["将交付 ZIP 解压到本地文件夹，不要直接在压缩包内运行。", "双击“系统缓存清理工具V1.0.exe”。", "如 Windows 显示安全提示，请核对文件名称和来源后继续。", "需要处理 Windows 临时目录时，可右键程序并选择“以管理员身份运行”。"]:
    bullet(doc, text)

doc.add_heading("3. 软件操作", 1)
doc.add_heading("3.1 首页", 2)
doc.add_paragraph("首页显示系统盘的总容量、已使用空间和可用空间。页面底部说明软件只处理预设缓存目录。")
add_picture(doc, "01_home.png", "图 1  软件首页（实际运行界面）")
doc.add_paragraph("操作：核对磁盘信息后点击“开始扫描”。扫描阶段仅统计文件，不会删除任何内容。")
doc.add_heading("3.2 扫描缓存", 2)
doc.add_paragraph("程序依次检查三类缓存位置，并实时显示当前类别、已发现文件数量和占用空间。扫描耗时取决于文件数量及磁盘速度。")
add_picture(doc, "02_scan.png", "图 2  缓存扫描界面（实际运行界面）")
doc.add_paragraph("操作：保持程序运行并等待扫描结束。扫描过程中不建议关闭窗口。")
doc.add_heading("3.3 选择清理项目", 2)
doc.add_paragraph("扫描完成后，列表按类别显示说明、文件数量及占用空间。用户可通过复选框调整清理范围。")
add_picture(doc, "03_results.png", "图 3  扫描结果界面（实际运行界面）")
doc.add_paragraph("操作：逐项核对结果，取消不需要清理的类别，然后点击“清理选中项目”。若未选中任何项目，软件不会进入清理流程。")
doc.add_heading("3.4 确认并执行清理", 2)
doc.add_paragraph("软件在删除前显示选中类别和预计释放空间。只有选择“是”后才执行清理；选择“否”将返回结果页。")
add_picture(doc, "04_clean.png", "图 4  清理执行界面（实际运行界面）")
doc.add_paragraph("操作：清理过程中保持程序运行。文件被占用、权限不足或路径校验失败时，程序记录失败并继续处理其他文件。")
doc.add_heading("3.5 查看清理结果", 2)
doc.add_paragraph("任务完成后显示实际释放空间、成功文件数、失败文件数、耗时及日志位置。")
add_picture(doc, "05_complete.png", "图 5  清理完成界面（实际运行界面）")
doc.add_paragraph("操作：记录结果后点击“返回首页”。如失败数量不为 0，可关闭相关应用或使用管理员权限重新扫描。")

doc.add_heading("4. 安全机制", 1)
for text in ["只扫描代码中明确配置的三类缓存目录。", "每个文件删除前再次校验其真实路径必须位于对应的允许目录内。", "拒绝处理符号链接、目录以及扫描后失效或被替换的路径。", "不提供任意目录输入，避免用户误选系统盘或个人目录。", "权限不足和文件占用按失败记录，不会导致整个任务中断。", "每次清理生成 UTF-8 文本日志，便于核对处理结果。"]:
    bullet(doc, text)
doc.add_paragraph("重要提示：缓存文件删除后通常无法恢复。执行清理前应保存正在编辑的工作，并关闭可能占用缓存的应用程序。")

doc.add_heading("5. 异常处理", 1)
add_table(doc, ["现象", "可能原因", "处理方法"], [("部分系统缓存无法扫描", "当前用户权限不足", "退出软件，右键选择“以管理员身份运行”"), ("失败文件数大于 0", "文件被系统或应用占用", "关闭相关应用后重新扫描；也可保留失败项"), ("显示 0 B 可清理", "缓存为空或目录不可访问", "无需清理；必要时检查运行权限"), ("窗口启动较慢", "单文件 EXE 首次解压运行环境", "等待数秒，避免连续重复双击"), ("安全软件提示未知程序", "个人项目未配置商业代码签名", "核对 GitHub 仓库和文件来源，在隔离环境中验证")], [1.5, 2.0, 2.9])

doc.add_heading("6. 日志说明", 1)
doc.add_paragraph("日志默认保存在当前用户目录下：AppData\\Local\\SystemCacheCleaner\\logs。日志记录开始时间、缓存类别、文件路径、成功或失败状态、释放字节数和耗时。日志可能包含本机临时文件路径，对外发送前应先检查内容。")

doc.add_heading("7. 验收检查", 1)
add_table(doc, ["编号", "检查项", "预期结果"], [("T01", "启动 EXE", "主界面正常显示，无命令行窗口"), ("T02", "点击开始扫描", "界面保持响应并最终显示三类结果"), ("T03", "取消一个类别后清理", "只处理仍被选中的类别"), ("T04", "确认窗口选择否", "不删除文件并返回结果页"), ("T05", "遇到无权限文件", "记录失败并继续，不崩溃"), ("T06", "完成清理", "显示空间、数量、耗时并生成日志")], [0.7, 2.4, 3.3])

doc.add_heading("8. 常见问题", 1)
for q, a in [("清理会删除个人文档吗？", "不会。V1.0 不扫描桌面、文档、下载、图片等个人目录。"), ("为什么每次扫描的空间不同？", "缓存会随系统和应用运行动态变化，这是正常现象。"), ("为什么预计空间与实际释放空间不同？", "扫描后文件可能被应用删除、占用或修改，实际结果以完成页和日志为准。"), ("没有安装 Python 可以运行吗？", "可以。Release 中的 EXE 已包含运行所需组件。")]:
    p = doc.add_paragraph(); set_font(p.add_run("问：" + q), 10, True); set_font(p.add_run("\n答：" + a), 10)

doc.add_heading("9. 版本记录", 1)
add_table(doc, ["版本", "日期", "变更内容"], [("V1.0", "2026-09-08", "完成缓存扫描、分类选择、安全清理、结果统计、日志和操作说明书")], [1.0, 1.4, 4.0])
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("—— 文档结束 ——"), 9, False, "64748B")
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
